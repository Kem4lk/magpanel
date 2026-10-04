# -*- coding: utf-8 -*-
"""
PCB uretim makinesi, prototip v0: Turkiye'den alinabilecek parcalar (BOM).
Tek kaynak bu dosya: satirlari burada duzelt, `python3 bom_v0.py` ile bom-v0.xlsx'i yeniden uret
(openpyxl gerekir: pip install openpyxl).

Fiyatlar 2026-10-04'te satici sayfasinda gorulen KDV dahil TL fiyatlaridir; stok ve fiyat degisir.
Dogrulandi = "evet": urun sayfasi acildi, fiyat/stok sayfanin kendi verisinden okundu.
"hayir": yalniz arama sonucu ozetinde goruldu ya da sayfa acilmadi (pazaryerleri Cloudflare ile engelliyor).
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DATE = '2026-10-04'
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'bom-v0.xlsx')

POZ, CNC, ISL, HOT, ELK = ('1 Pozlama', '2 CNC portal', '3 Islak istasyon', '4 Hot plate', '5 Elektronik')

# (modul, kalem, ozellik, adet, birim, birim_fiyat|None, satici, url, stok, dogrulandi, not)
BOM = [
    # --- 1. UV LCD maske pozlama unitesi ---
    (POZ, 'Mono LCD 4K', 'Phrozen Sonic Mighty 4K, 9.3", 3840x2400, ~52 µm, ~200x125 mm; MIPI FPC (HDMI yok)', 1, 'adet', 11280.00,
     '3Dream', 'https://store.3dream.com.tr/products/phrozen-lcd-ekran', 'Stokta', 'evet',
     'ESP32-P4 için en uygun (kare ~9 MB). Ucuz alt.: Elegoo Saturn 2/8K 10" 7680x4320, 4.700 TL, 3Dream, stokta; 8K kare P4 PSRAM\'ine sığmaz. Init dizisi yayımlanmamış, deneyle çıkarılacak.'),
    (POZ, 'UV LED ışık modülü', 'Anycubic Photon Mono M7 Pro ışık kartı, 405 nm matris, 10" sınıfı', 1, 'adet', 4230.00,
     '3Dream', 'https://store.3dream.com.tr/products/anycubic-uv-led-modulu', 'Stokta', 'evet',
     'Gerilim/akım sayfada yok: kart etiketinden ölçülmeli. Alt.: Creality Halot Lite/Sky UV light, Robo90 470,40 TL (fiyat şüpheli düşük, teyit et). Gerçek 405 nm COB TR\'de yok → Yurt dışı sayfası.'),
    (POZ, 'Fresnel mercek', 'Anycubic Mono M7 / M7 Pro Fresnel (ışık modülüyle aynı set)', 1, 'adet', 1645.00,
     '3Dream', 'https://store.3dream.com.tr/products/anycubic-fresnel-lens', 'Stokta', 'evet',
     'Alt.: A4 sayfa büyüteci 180x120 PVC, Hepsiburada ~398 TL (doğrulanamadı; odak ve 405 nm geçirgenliği belirsiz).'),
    (POZ, 'LCD koruma filmi', 'Anycubic M7/M7 Pro ekran koruyucu, 5\'li', 1, 'paket', 1176.00,
     'Robo90', 'https://www.robo90.com/anycubic-ekran-koruyucu-5-adet-mono-m7m7-pro', 'Stokta', 'evet',
     'Panel ile kart arası kalınlık bulanıklık yapar: ince koruyucu tercih et. Phrozen 4K\'ya ölçü uyumu kontrol edilmeli.'),
    (POZ, 'Sabit akım LED sürücü', 'XL4015 ekranlı DC-DC, CC+CV, giriş ≤38 V, 5 A', 1, 'adet', 418.45,
     'Direnc.net', 'https://www.direnc.net/xl4015-display-dc-dc-power-modul', 'Stokta', 'evet',
     'Direnc\'in 98,80 TL\'lik "XL4015 5A Ayarlanabilir" modülü sabit akım DEĞİL. Ürün için Mean Well LDD-H (PWM kısılabilir) → Yurt dışı.'),
    (POZ, 'Soğutucu + fan', '80x80x54 mm alüminyum işlemci soğutucusu, 12 V 0,15 A fan', 1, 'adet', 390.00,
     'Motorobit', 'https://www.motorobit.com/80x80x54mm-sogutuculu-fan-islemci-fani', 'Stokta', 'evet',
     'Braket yok. Işık modülü büyükse alt.: 200x195x10 mm alüminyum plaka, Motorobit 1.200 TL + fan.'),
    (POZ, 'ESP32-P4 geliştirme kartı', 'Waveshare ESP32-P4-NANO: MIPI-DSI/CSI, 32 MB PSRAM, 16 MB flash, C6 ile Wi-Fi 6', 1, 'adet', 2400.00,
     'Motorobit', 'https://www.motorobit.com/esp32-p4-nano-yuksek-performansli-gelistirme-karti', 'Stokta', 'evet',
     'Alt.: Waveshare ESP32-P4-Pico, Robotistan 1.646,73 TL, stokta. Espressif Function-EV-Board TR\'de yok.'),
    (POZ, 'Sıcaklık sensörü (LED)', 'DS18B20, M10 vidalı prob, 1 m', 1, 'adet', 312.00,
     'Motorobit', 'https://www.motorobit.com/ds18b20-m10-dijital-sicaklik-sensoru-1-metre', 'Stokta', 'evet',
     'Ucuz alt.: çıplak DS18B20 36 TL (Motorobit) ya da 10k NTC 6,74 TL (Direnc).'),
    (POZ, 'UV koruma gözlüğü', 'Univet 5X7.03.00.04, turuncu lens, ~525 nm\'ye kadar keser', 1, 'adet', 1489.08,
     'Koçtaş', 'https://www.koctas.com.tr/univet-5x7030004-turuncu-lensli-gozluk/p/5000208517', '?', 'hayır',
     'LED pozlama için yeterli. Profesyonel alt.: GHP 180–540 nm OD5+, lazergozluk.net 17.650 TL.'),
    (POZ, 'UV engelleyici kapak malzemesi', '3 mm turuncu akrilik (pleksi), ölçüye kesim', 1, 'levha', None,
     'Hepsiburada (TeknoTrust)', 'https://www.hepsiburada.com/teknotrust-3-mm-akrilik-pleksi-levha-seffaf-kalinlik-tum-ebatlar-icin-gecerlidir-pm-HBC00008CFKNV', '?', 'hayır',
     'Sıradan turuncu pleksinin 405 nm\'yi kestiği garanti değil, test et. Asıl yol: reçine yazıcı kapağı / UV engelleyici akrilik → Yurt dışı.'),

    # --- 2. CNC delme / kenar frezeleme portali (ozel yapim, B) ---
    (CNC, 'Lineer ray HGR15', '15 mm, Hiwin uyumlu, metre fiyatı, ücretsiz kesim', 2, 'm', 766.51,
     'cnc-marketi', 'https://www.cnc-marketi.com/urun/lineer-kizak-15-mm-lineer-kizak-fiyatlari', 'Stokta (572 m)', 'evet',
     'X/Y/Z ikişer ray ~300 mm. Alt.: Rulmansepetim 1.236,27 TL/m. Daha hafif: MGN12 400 mm set, Motorobit 1.080 TL.'),
    (CNC, 'Lineer araba HGH15CA', 'Dar tip, Hiwin uyumlu', 8, 'adet', 389.96,
     'cnc-marketi', 'https://www.cnc-marketi.com/urun/15-mm-dar-lineer-kizak-arabasi-17737367412859', 'Stokta', 'evet',
     'Alt.: Rulmansepetim 431,41 TL.'),
    (CNC, 'Bilyalı vida SFU1204', 'Uçları tornalanmış (BK/BF10), somun dahil, 350 mm', 3, 'adet', 2525.52,
     'ileri3d', 'https://www.ileri3d.com/urun/sfu-1204-tornalanmis-mil-ve-vidali-somun-350mm', 'Stokta görünüyor', 'evet',
     '250/300 mm tükendi. Haddelenmiş C7 tek somun birkaç 10 µm boşluk taşır: ±0,02 mm için ön yüklü/C5 gerekebilir (Mermak/Rulmansepetim\'e sor).'),
    (CNC, 'BK10 sabit yatak', '1204 için', 3, 'adet', 313.32,
     'cnc-marketi', 'https://www.cnc-marketi.com/urun/bk-10-vidali-mil-uc-yataklari', 'Stokta', 'evet', 'Alt.: ileri3d 1.036,11 TL.'),
    (CNC, 'BF10 serbest yatak', '8 mm mil', 3, 'adet', 201.71,
     'cnc-marketi', 'https://www.cnc-marketi.com/urun/bf-10-vidali-mil-uc-yataklari-8mm', 'Stokta', 'evet', 'Alt.: ileri3d 518,05 TL.'),
    (CNC, 'Somun gövdesi 1204', '', 3, 'adet', 548.67,
     'ileri3d', 'https://www.ileri3d.com/urun/somun-govdesi', 'Stokta görünüyor', 'evet',
     'Alt.: Rhino3D 326,66 TL.'),
    (CNC, 'Kaplin 5x8', 'Alüminyum sabit (rijit)', 3, 'adet', 45.00,
     'Motorobit', 'https://www.motorobit.com/aluminyum-sabit-kaplin-5x8mm', 'Stokta (100)', 'evet',
     'Esnek/çeneli tercih edilir: cnc-marketi GS14 409,60 TL (5x8 için sor).'),
    (CNC, 'Step motor NEMA17', '0,47 Nm', 3, 'adet', 471.26,
     'cnc-marketi', 'https://www.cnc-marketi.com/urun/0-47-nm-step-motor-nema-17', 'Stokta', 'evet',
     'Alt.: NEMA23 2,2 Nm, Motorobit 1.800 TL.'),
    (CNC, 'Step sürücü TMC5160', 'MKS TMC5160 PRO, SPI, dahili rampa üreteci', 3, 'adet', 1606.86,
     'Rhino3D Printer', 'https://www.rhino3dprinter.com/urun/mks-tmc5160-pro-step-motor-surucusu', 'Stokta', 'evet',
     'Ucuz alt.: TMC2209 V2.0, Motorobit 390 TL (≤29 V, rampa üreteci yok).'),
    (CNC, 'ER11 mil seti', '400 W, 12–48 V, 3–12k rpm, φ52, ER11 + braket + hız sürücüsü', 1, 'set', 6900.00,
     'F1Depo', 'https://www.f1depo.com/urun/0-4-kw-spindle-motor-er11-bracket-tutma-aparati', 'Belirsiz ("Gelince haber ver")', 'evet',
     'Stok teyit et. Alt.: Motorobit 200 W 12–48 V mil 2.220 TL + ER11 şaft tutucu 510 TL. Salgı (runout) ölçülmeli.'),
    (CNC, 'Mil güç kaynağı 48 V', '48 V 7,3 A, 350 W', 1, 'adet', 1859.76,
     'Direnc.net', 'https://www.direnc.net/24v-165a-metal-switch-mod-adaptor', 'Stokta', 'TEYİT ET',
     'Link adı "24V 16.5A" gösteriyor: 48 V model olduğunu siparişten önce teyit et.'),
    (CNC, 'ER11 pens 3,175 mm (1/8")', 'PCB matkap/freze sapı için', 2, 'adet', 150.00,
     'Motorobit', 'https://www.motorobit.com/er11-3175mm-pens-18', 'Stokta (6)', 'evet',
     'cnc-marketi 7\'li sette 3,175 YOK.'),
    (CNC, 'PCB matkap seti', '0,3–1,2 mm, 10 adet, 3,175 mm sap', 2, 'set', 270.00,
     'Motorobit', 'https://www.motorobit.com/pcb-matkap-ucu-seti-03-1', 'Stokta (31)', 'evet',
     '3 mm pim delikleri için ayrı 3,0 mm karbür matkap gerekir (aranmadı).'),
    (CNC, 'Kenar frezesi', 'Karbür "elmas/mısır koçanı" 1,6 mm, 3,17 mm sap', 3, 'adet', 150.00,
     'Motorobit', 'https://www.motorobit.com/elmas-freze-matkap-ucu-160mm-x-11mm', 'Stokta', 'evet',
     '1,85/2,05/2,45 mm de var.'),
    (CNC, 'Sigma profil 40x40', '8 kanal hafif, ücretsiz kesim', 3, 'm', 774.84,
     'cnc-marketi', 'https://www.cnc-marketi.com/urun/sigma-profil-40x40-sigma-profil-8-kanal', 'Stokta (289 m)', 'evet',
     'Alüminyum plaka (tabla, Z plakası) için yerel lazer/CNC kesimciden teklif.'),
    (CNC, 'Köşe bağlantı 40x40 K8', 'Geniş köşe', 16, 'adet', 38.51,
     'cnc-marketi', 'https://www.cnc-marketi.com/urun/genis-kose-baglanti-40x40-k8', 'Stokta', 'evet', ''),
    (CNC, 'T-somun M5 K8', '', 100, 'adet', 3.67,
     'cnc-marketi', 'https://www.cnc-marketi.com/urun/t-kanal-somunu-m5-30x30-k8', 'Stokta', 'evet', ''),
    (CNC, 'Endüktif limit sensörü', 'LJ12A3-4-Z/BX, NPN', 3, 'adet', 116.24,
     'Direnc.net', 'https://www.direnc.net/enduktif-npn-proximity-sensor-lj12a3-4-zbx', 'Stokta (152)', 'evet',
     'Alt.: mekanik ME-8111, Motorobit 126 TL.'),
    (CNC, 'Takım sıfırlama probu', 'Yaylı Z probu', 1, 'adet', 2358.33,
     'Meon Otomasyon', 'https://www.meonotomasyon.com/urun/cnc-takim-sifirlama-probu', 'Muhtemelen stokta', 'evet',
     'Yükseklik haritası için bakır–uç süreklilik probu (iki kablo) da yeter.'),
    (CNC, 'Konum pimi 3 mm', 'DIN 6325 / ISO 8734, sertleştirilmiş, 3x20 mm', 10, 'adet', None,
     'RS Components TR', 'https://tr.rsdelivers.com/product/rs-pro//rs-pro-duz-tespit-pimi-cap-3mm-20mm-celik/0270568', 'Teklif', 'hayır',
     'Fiyat görülmedi. Alt.: Norelem/KIPP TR teklif, Amazon.com.tr ilanları (açılmadı).'),

    # --- 3. Islak istasyon + kimyasallar + sarf ---
    (ISL, 'PP kapaklı tank', 'IKEA 365+ 4,2 L, 21x15x23 cm, kap+kapak PP', 4, 'adet', 329.00,
     'IKEA', 'https://www.ikea.com.tr/urun/ikea-365--plastik-4-2-lt-kapakli-plastik-saklama-kabi-59276822', 'Stokta', 'evet',
     'Banyo / aşındırma / sökme / durulama. 100 mm kart dik sığar.'),
    (ISL, 'Isıtıcı', 'Newa Therm 50 W, cam tüp, termostatlı', 1, 'adet', 1554.08,
     'Evcilal', 'https://www.evcilal.com/urun/newa-therm-50-w', 'Stokta', 'evet',
     'Akvaryum termostatı ~32–34 °C\'de keser (tipik): 40–50 °C için ESP32 + SSR ile sür. Titanyum küçük ısıtıcı TR\'de yok.'),
    (ISL, 'Hava pompası', 'Sobo SB-666A, çift çıkış, 5 W, 2x4 L/dk', 1, 'adet', 375.00,
     'Mixpet', 'https://www.mixpet.com.tr/sobo-cift-cikisli-akvaryum-hava-motoru-5w-2x4lt/dak-sb-666a', 'Stokta', 'evet',
     'Pompayı sıvı seviyesinin üstüne koy.'),
    (ISL, 'Silikon hava hortumu', '4x6 mm, 5 m', 1, 'adet', 75.00,
     'Badem Akvaryum', 'https://www.bademakvaryum.com/akvaryum-hava-hortumu-silikonlu-5-metre-1199', 'Stokta', 'evet', ''),
    (ISL, 'Çekvalf', '4/6 mm', 2, 'adet', 30.00,
     'Badem Akvaryum', 'https://www.bademakvaryum.com/akvaryum-cek-valf-hava-hortum-aparati-385', 'Stokta', 'evet',
     'Şart: asit pompaya geri emilmesin.'),
    (ISL, 'Hava taşı', 'Düz, 15 cm', 1, 'adet', 70.00,
     'Badem Akvaryum', 'https://www.bademakvaryum.com/akvaryum-hava-tasi-duz-15-cm-630', 'Stokta', 'evet',
     'Seramik/mineral taş asitte dağılabilir: en sağlamı delikli PP boru.'),
    (ISL, 'ORP elektrodu (BNC)', 'SEKO, ±1000 mV, 60 °C, epoksi gövde, 1,5 m kablo', 1, 'adet', 8146.67,
     'Havuzdan.com', 'https://www.havuzdan.com/urun/seko-dozaj-pompasi-orp-elektrodu-1-5-mt-kabloluwf-sk', 'Sayfada yazmıyor', 'evet',
     'Alt.: ISOLAB ORP probu 7.225,44 TL (interlab); EMEC ERHL 9.576 TL (sentezlab). Ucuz BNC problar yurt dışından.'),
    (ISL, 'ORP için ADC', 'ADS1115 16-bit 4 kanal ADC (+ op-amp tampon devresi kendin)', 1, 'adet', 134.44,
     'Robotistan', 'https://www.robotistan.com/ads1115-16-bit-4-kanal-adc', 'Stokta (157)', 'evet',
     'Hazır DFRobot SEN0165 ORP kartı TR\'de yok → Yurt dışı.'),
    (ISL, 'Peristaltik pompa', '12 V, PharMed BPT hortum, ≥47 ml/dk (NKP-DC-B08D)', 2, 'adet', 553.31,
     'Robotistan', 'https://www.robotistan.com/12-v-peristaltik-sivi-pompasi-bpt-tup-nkp-dc-b08d', 'Stokta (28)', 'evet',
     'HCl/H2O2 için BPT hortum şart (silikon değil). Yedek hortum al. Alt.: Kamoer B08G, Motorobit 690 TL.'),
    (ISL, 'DS18B20 su geçirmez', '6x30 mm paslanmaz kılıf, 1 m', 2, 'adet', 75.89,
     'Robotistan', 'https://www.robotistan.com/su-gecirmez-ds18b20-dijital-isi-sensoru', 'Stokta (213)', 'evet',
     'Paslanmaz kılıf CuCl2/HCl\'de aşınır: PTFE makaron ya da cam kılıfa al.'),
    (ISL, 'Tuz ruhu (HCl)', '1 kg (yüzde sayfada yok)', 2, 'kg', 100.00,
     'ehammaddem', 'https://ehammaddem.com/products/hidroklorik-asit-tuz-ruhu-1-kg', 'Stokta', 'evet',
     'Yoğunluğu ölç. Alt.: Kimyacınız %33 125 TL (yalnız mağazadan teslim).'),
    (ISL, 'Hidrojen peroksit %50', 'Perhidrol, 1 L', 1, 'L', 200.00,
     'ehammaddem', 'https://ehammaddem.com/products/hidrojen-peroksit-perhidrol-50-1-litre', 'Stokta', 'evet',
     'ADR sınıf 5.1: kargo reddedebilir. Alt.: Kimyacınız %50 1 kg 175 TL.'),
    (ISL, 'Kostik (NaOH)', 'Boncuk, 1 kg', 1, 'kg', 180.00,
     'ehammaddem', 'https://ehammaddem.com/products/sodyum-hidroksit-boncuk-kostik-1-kg', 'Stokta', 'evet', 'Pozitif lak banyosu 7 g/L.'),
    (ISL, 'Sodyum karbonat', 'Toz çamaşır sodası, 1 kg', 1, 'kg', 75.00,
     'ehammaddem', 'https://ehammaddem.com/products/sodyum-karbonat-toz-camasir-sodasi-1-kg', 'Stokta', 'evet', 'Dry film banyosu %1.'),
    (ISL, 'Aseton', '1 L', 1, 'L', 200.00,
     'Kimyacınız', 'https://www.kimyaciniz.com/aseton-1-lt', 'Stokta', 'evet',
     'Sayfa "şahıslara satış yapılmaz" diyor; hırdavat/eczaneden de alınır.'),
    (ISL, 'Distile su', '0,00–0,02 ppm, 5 L', 1, 'adet', 400.00,
     'Kimyacınız', 'https://www.kimyaciniz.com/distile-saf-su-0-00-0-02-ppm-5-lt', 'Stokta', 'evet',
     'Boş bidon atık kabı olarak kullanılabilir.'),
    (ISL, 'Positiv 20', 'Kontakt Chemie, 200 ml sprey', 1, 'adet', 1224.65,
     'Motorobit', 'https://www.motorobit.com/positiv-20-baski-devre-pozlama-spreyi-200ml', 'Stokta (119)', 'evet', ''),
    (ISL, 'Bakır plaket tek yüz', 'FR4 1,6 mm, 15x20 cm (iki adet 15x10 çıkar)', 3, 'adet', 87.24,
     'Robotistan', 'https://www.robotistan.com/20x20-bakir-plaket-fr4', 'Stokta (235)', 'evet', ''),
    (ISL, 'Bakır plaket çift yüz', 'FR4 1,6 mm, 35 µm, 15x20 cm', 3, 'adet', 87.24,
     'Robotistan', 'https://www.robotistan.com/20x20-cift-yuzlu-bakir-plaket-fr4-epoksi', 'Stokta (192)', 'evet', ''),
    (ISL, 'Nitril kimyasal eldiven', 'Starline STL-1513, 330 mm, 0,38 mm, EN 374', 2, 'çift', 125.00,
     'Entegre Safety', 'https://www.entegresafety.com/urun/nitril-eldiven-22', 'Aynı gün kargo', 'evet',
     'Nitril asetona karşı zayıf.'),
    (ISL, 'Kimyasal gözlük (goggle)', '3M 2890, dolaylı havalandırma, EN 166 "3"', 1, 'adet', 985.83,
     'Entegre Safety', 'https://www.entegresafety.com/urun/3m-2890-goggle-koruyucu-gozluk-clear-as-af', '"Gelince haber ver"', 'evet',
     'Alt.: 3M 2890A 1.167,96 TL aynı satıcı.'),
    (ISL, 'Kimyasal önlük', 'Borox PVC, 350 µm, 70x110 cm', 1, 'adet', 448.14,
     'Blabmarket', 'https://www.blabmarket.com/urun/asit-ve-kimyasal-onlugu-kimyasallara-dayanikli-pvc-onluk-su-gecrmez-is-guvenligi-onlugu-unisex-tam-boy-onluk', 'Stokta', 'evet', ''),
    (ISL, 'Atık bidonu', 'ISOLAB HDPE 3 L, geniş boyun', 1, 'adet', None,
     'Derkimlab', 'https://www.derkimlab.com/urun/isolab-bidon-hdpe-3-lt', '3 günde kargo', 'hayır',
     'Fiyat gösterilmiyor. Bakırlı atık kanalizasyona dökülmez.'),

    # --- 4. Reflow hot plate (kendin yap, urun yolu) ---
    (HOT, 'Isıtıcı plaka', 'PTC alüminyum 220 V 600 W ~260 °C (ya da silikon ped + 6 mm Al plaka)', 1, 'adet', None,
     'Amazon.com.tr (Hyuduo)', 'https://www.amazon.com.tr/Hyuduo-istasyonu-lehimleme-plakas%C4%B1-%C3%A7%C4%B1kar%C4%B1c%C4%B1/dp/B09P3RP4ZJ', '?', 'hayır',
     'Elektronik mağazalarında yok. Ürün için ~200x150 mm 500 W döküm/plaka rezistans teklifi al. Hazır alt.: Sunline 958 130x130, Robotistan 13.145,64 TL (Alternatifler).'),
    (HOT, 'Termokupl okuyucu', 'MAX6675 + K-tip prob, SPI', 1, 'adet', 293.39,
     'Robotistan', 'https://www.robotistan.com/max6675-k-type-termokupl-sensor', 'Stokta (149)', 'evet', ''),
    (HOT, 'K-tip prob (vidalı)', 'M6 vidalı, 0–800 °C, 50 cm', 1, 'adet', 150.00,
     'Motorobit', 'https://www.motorobit.com/k-tip-0-800c-m6-termokupl-50cm', 'Stokta (55)', 'evet', 'Plakaya vidalanır.'),
    (HOT, 'SSR', 'SSR-40DA, 3–32 VDC giriş, 24–380 VAC', 2, 'adet', 215.03,
     'Direnc.net', 'https://www.direnc.net/ssr-40da--40a-solid-state-role', 'Stokta (876)', 'evet',
     'Biri hot plate, biri ıslak istasyon ısıtıcısı. Fotek tipi klonlar şişirilmiş değerli: ≥2x akım + soğutucu.'),
    (HOT, 'SSR soğutucu', 'I-50, 10–40 A', 2, 'adet', 150.00,
     'Motorobit', 'https://www.motorobit.com/i-50-ssr-10-40a-role-uyumlu-sogutucu', 'Stokta (28)', 'evet', ''),
    (HOT, 'Krem lehim', 'Sn63Pb37, 50 g', 1, 'adet', 333.00,
     'Motorobit', 'https://www.motorobit.com/tassol-krem-lehim-50gr-sn63-pb37', 'Stokta', 'evet',
     'TR\'de "lehim pastası" = flux; lehim için "krem lehim" ara. Kurşunsuz alt.: Tassol 486 TL (PTC plaka SAC305 için sınırda).'),
    (HOT, 'Flux', 'Kurtel RA 50 reçineli, 10 cc', 1, 'adet', 443.60,
     'Robotistan', 'https://www.robotistan.com/kurtel-ra-50-recineli-flux-kalemi-10cc', 'Stokta (4)', 'evet', ''),
    (HOT, 'Duman emici', 'FA-400 ESD + 5 aktif karbon filtre', 1, 'adet', 1627.29,
     'Direnc.net', 'https://www.direnc.net/fa-400-esd-guvenli-duman-emici-ve-5-adet-aktif-karbon-filtre', 'Stokta (18)', 'evet', ''),

    # --- 5. Ortak elektronik ---
    (ELK, 'ESP32-S3 kart', 'ESP32-S3-WROOM-1-N16R8, çift USB-C (DevKitC-1 klonu)', 2, 'adet', 610.23,
     'Direnc.net', 'https://www.direnc.net/esp32-s3-n16r8-wifi-bluetooth-board', 'Stokta (1792)', 'evet',
     'CNC hareket + ıslak istasyon/hot plate denetimi.'),
    (ELK, '24 V güç kaynağı', 'Hightek HT-1212, 24 V 15 A 360 W, metal kasa', 1, 'adet', 1320.00,
     'Motorobit', 'https://www.motorobit.com/hightek-24v-15a-metal-kasa-adaptor-led-driver', 'Stokta (4)', 'evet',
     'Ürün için Mean Well: RSP-320-24 Direnc 3.635,83 TL stokta; LRS-350-24 Akakçe ~1.338 TL (doğrulanamadı).'),
    (ELK, '12 V güç kaynağı', 'Mean Well LRS-100-12, 12 V 8,5 A', 1, 'adet', 1009.85,
     'Direnc.net', 'https://www.direnc.net/12v-85a-kapali-tip-switch-mode-adaptor', 'Stokta', 'evet',
     'Pompalar, fanlar, LED sürücü girişi (LED 30+ V ise ayrı kaynak gerekir).'),
    (ELK, '5 V düşürücü', 'HCW-P715, 9–36 V → 5 V 5 A', 2, 'adet', 165.00,
     'Motorobit', 'https://www.motorobit.com/hcw-p715-9-36vtan-5va-5a-voltaj-dusurucu-regulator-karti', 'Stokta (75)', 'evet', ''),
    (ELK, 'MOSFET anahtarlama kartı', '15 A / 400 W PWM, tetik 3,3–20 V', 4, 'adet', 93.21,
     'Robotistan', 'https://www.robotistan.com/15a-400w-pwm-kontrollu-mosfet-anahtarlama-karti', 'Stokta (174)', 'evet',
     'Pompa, fan, hava pompası, LED. IRF520 modülleri lojik seviye değil.'),
    (ELK, '4 kanal röle', '5 V', 1, 'adet', 127.86,
     'Direnc.net', 'https://www.direnc.net/5v-4-kanal-role-karti', 'Stokta (691)', 'evet', ''),
    (ELK, 'Dokunmatik ekran (ops.)', '2,8" SPI TFT ILI9341 240x320', 1, 'adet', 569.45,
     'Robotistan', 'https://www.robotistan.com/28-inc-arduino-dokunmatik-ekran-shield-240320', 'Stokta (774)', 'evet',
     'Web arayüzü varken isteğe bağlı.'),
    (ELK, 'GX16 konnektör', '4 pin fiş+soket', 6, 'adet', 102.00,
     'Motorobit', 'https://www.motorobit.com/gx-16-4-pin-su-gecirmez-mike-konnektor', 'Stokta (1781)', 'evet', 'Modüller arası kablolar.'),
    (ELK, 'JST-XH kablo seti', '2,54 mm', 1, 'set', 372.26,
     'Robotistan', 'https://www.robotistan.com/xh-2-54mm-konnektor-kablo-seti', 'Stokta (31)', 'evet', ''),
    (ELK, 'Silikon kablo 18 AWG', 'siyah/kırmızı', 4, 'm', 42.00,
     'Motorobit', 'https://www.motorobit.com/18-awg-silikon-kablo-siyah-1-metre', 'Stokta (1691)', 'evet', ''),
    (ELK, 'Silikon kablo 20 AWG seti', '6 renk, 60 m', 1, 'set', 1440.00,
     'Motorobit', 'https://www.motorobit.com/20awg-silikon-cok-damarli-montaj-kablosu-seti-6-renk-60m', 'Stokta (69)', 'evet',
     '22 AWG silikon stokta yok.'),
    (ELK, 'DC jak', 'DC-099 5,5x2,1 mm metal panel', 2, 'adet', 18.00,
     'Motorobit', 'https://www.motorobit.com/dc-099-55x21mm-metal-dc-jack-sasesi-jak-girisi', 'Stokta (800)', 'evet', ''),
    (ELK, 'Sigorta yuvası', '5x20 panel tipi', 2, 'adet', 9.00,
     'Motorobit', 'https://www.motorobit.com/5x20-panel-tipi-sigorta-yuvasi', 'Stokta', 'evet', ''),
    (ELK, 'Sigorta T3,15 A', '5x20 gecikmeli cam', 10, 'adet', 3.60,
     'Motorobit', 'https://www.motorobit.com/5k-t315a-5x20mm-gecikmeli-cam-sigorta', 'Stokta', 'evet', ''),
    (ELK, 'Acil stop', 'Mantar başlı 1NO+1NC, anahtarlı reset', 1, 'adet', 1200.00,
     'Motorobit', 'https://www.motorobit.com/mantar-basli-acil-stop-butonu-1no1nc-anahtarli-lift-reset', 'Stokta (20)', 'evet',
     'Ucuz alt.: Direnc Drn956 16 mm 188,30 TL. Ürün için endüstriyel NC tercih.'),
    (ELK, 'IEC C14 giriş', 'EMI filtreli, 6 A, anahtarlı, sigorta yuvalı', 1, 'adet', 690.00,
     'Motorobit', 'https://www.motorobit.com/emi-filtre-6a-anahtarli-erkek-power-soketi-sigorta-yuvali', 'Stokta (95)', 'evet', ''),
]

# Ana listeye girmeyen secenekler: (modul, kalem, ozellik, fiyat|None, satici, url, durum, not)
ALT = [
    (CNC, 'Hazır CNC3018 kit (lazerli)', '300x180x45 mm, 775 mil (ER11 değil), düz mil + T-vida, GRBL', 26700.00,
     'Motorobit', 'https://www.motorobit.com/cnc3018-15000mw-lazerli-cnc-makinesi-kesim-tezga', 'Stokta (8)',
     'v0 hızlı deneme. ±0,1 mm: ±0,02 hedefini tutmaz. Aynı kit Robotistan 28.307,95 TL.'),
    (CNC, 'Hazır CNC3018 PRO MAX', '500 W φ52 mil 48 V, limit switch + Z probu + acil stop dahil', 50100.00,
     'Motorobit', 'https://www.motorobit.com/cnc3018-pro-max-15000mw-lazerli-cnc-makinesi-kesim-tezgahi', 'Stokta (5)',
     'v0 için daha uygun kit; sürücüler A4988, T-vida.'),
    (CNC, 'Hazır DRNC 3025', '30x25 cm, bilyalı vida + 15 mm ray, 600 W ER11 mil', 104611.61,
     'Direnc.net', 'https://www.direnc.net/drnc-3025-masaustu-cnc-ve-lazer-makinesi', 'Stokta (1)',
     'Hedef mimariye en yakın hazır makine.'),
    (CNC, 'CNC3018 lazersiz (en ucuz)', 'Temel 3018', 9899.86,
     'Robotistan', 'https://www.robotistan.com/cnc3018-masaustu-cnc-makinesi', 'Stokta yok', ''),
    (CNC, 'MGN12 ray + MGN12H', '400 mm set', 1080.00,
     'Motorobit', 'https://www.motorobit.com/mgn12-lineer-ray-kizak-400mm-mgn12h-bilyali-rulm', 'Stokta (14)', 'HGR15 yerine hafif seçenek.'),
    (CNC, 'SFU1605 vidalı mil', 'SCR1605, metre fiyatı, uç işleme yok', 1575.35,
     'cnc-marketi', 'https://www.cnc-marketi.com/urun/vidali-mil-1605-vidali-mil', 'Stokta (267 m)', 'Somun ayrı 573,02 TL (stokta yok).'),
    (CNC, 'TMC2209 V2.0', '1,7 A RMS, 4,75–29 V', 390.00,
     'Motorobit', 'https://www.motorobit.com/tmc2209-v20-step-motor-surucusu', 'Stokta (25)', 'TMC5160 yerine ucuz (3 adette ~3.650 TL düşer).'),
    (POZ, 'Elegoo Saturn 2 / 8K LCD', '10", 7680x4320, ~28,5 µm', 4700.00,
     '3Dream', 'https://store.3dream.com.tr/products/elegoo-lcd-ekran-saturn-8k-10-8k-saturn-2', 'Stokta',
     '150x100\'ü kapsayan en ucuz mono LCD; ESP32-P4 ile sürmek zor.'),
    (POZ, 'Anycubic M7 Pro anakartı', 'Yazıcının kendi LCD/ışık denetleyicisi', 7050.00,
     '3Dream', 'https://store.3dream.com.tr/collections/yedek-parca-anycubic', 'Stokta',
     'İlk deneyde M7 Pro LCD + ışık + Fresnel\'i kendi kartıyla sürme yolu (kapalı yazılım, yalnız deney).'),
    (POZ, 'Waveshare ESP32-P4-Pico', '32 MB PSRAM, MIPI DSI/CSI', 1646.73,
     'Robotistan', 'https://www.robotistan.com/waveshare-esp32p4pico-riscv-gelistirme-karti', 'Stokta', ''),
    (POZ, '3 W UV power LED', '395–400 nm, 600–700 mA', 87.95,
     'Roboshop', 'https://www.roboshop.com.tr/3w-power-led-uv-mor-395-400nm', 'Stokta',
     'Kendin-yap dizi (10–16 adet). 405 değil ama Positiv 20 için uygun.'),
    (HOT, 'Hazır ön ısıtıcı Sunline 958', '130x130 mm, 600 W, 50–400 °C', 13145.64,
     'Robotistan', 'https://www.robotistan.com/sunline-958-havya-istasyonu', 'Stokta (5)',
     '150x100 kart 20 mm taşar. Direnc 13.521,05 TL.'),
    (HOT, 'Nofaner 200x200 hot plate', '220 V, dijital', None,
     'Amazon.com.tr', 'https://www.amazon.com.tr/Nofaner-istasyonu-elektrikli-lehimleme-d%C3%BCzenlemesi/dp/B0B7BG89KB', '?',
     'TR\'de bulunan tek 200x200; doğrulanamadı.'),
    (ISL, 'Perklorür (FeCl3) sıvı', '1 kg', 125.00,
     'Kimyacınız', 'https://www.kimyaciniz.com/demir-3-klorur-sivi-1-kg', 'Stokta', 'Kupri klorür yerine tek kullanımlık aşındırıcı.'),
    (ISL, 'Sodyum persülfat', '1 kg', 584.67,
     'Kimya Depom', 'https://www.kimyadepom.com/urun/sodyum-persulfat-25-kg-torba', 'Sipariş verilebilir', 'Temiz, şeffaf aşındırıcı.'),
    (ISL, 'ISOLAB ORP probu', 'Ölçüm cihazları için', 7225.44,
     'Interlab', 'https://shop.interlab.com.tr/urun/orp-probu-olcum-cihazlari-icin', '?', 'SEKO\'dan ucuz.'),
    (ISL, 'UV lehim maskesi (yeşil)', 'Mechanic MT-UVH900, 10 cc', None,
     'Hepsiburada', 'https://www.hepsiburada.com/mechanic-mt-uvh900-ly-uv-isikla-sertlesen-solder-mask-devre-yalitim-boyasi-lehim-maskesi-10cc-yesil-pm-HBC0000FVNW47', '?',
     'v2 lehim maskesi pozlaması için.'),
    (ISL, 'Dry film (negatif)', '30 cm x 5 m', None,
     'n11', 'https://www.n11.com/urun/devre-uretim-icin-kuru-film-fotoresist-carsaf-duzleme-islemi-30cm-x-5m-pcb-podiksiyatler-kuru-film-dogruluk-deneyimli-film-79346576', '?',
     'Arama özetinde ~1.191 TL; çoğu yurt dışından.'),
]

# Turkiye'de bulunamayanlar
YURTDISI = [
    (POZ, 'HDMI sürücü kartlı mono LCD kiti (2K/4K)', 'İlk ışık motoru deneyi; panel arayüzünü çözmeden pozlama.',
     'AliExpress: "mono LCD 4K HDMI driver board" (mono olduğunu kontrol et, çoğu ilan RGB).'),
    (POZ, '405 nm COB LED 10–50 W', 'Kendi ışık motoru tasarımı.', 'AliExpress / LED üreticisi; TR\'de yalnız 395–400 nm 3 W LED var.'),
    (POZ, 'Mean Well LDD-1000H / LDD-1500H', 'PWM ile kısılabilir sabit akım LED sürücü (ürün sınıfı).', 'Mouser / Digikey / AliExpress.'),
    (POZ, 'UV engelleyici turuncu akrilik / reçine yazıcı kapağı', 'Kabin penceresi, 405 nm kesmesi garanti.', 'AliExpress / yazıcı yedek parçası.'),
    (ISL, 'DFRobot Gravity SEN0165 ORP kartı', 'Hazır ORP yükselticisi (±2000 mV).', 'AB ~€98 (opencircuit.shop) ya da DFRobot; TR alternatifi ADS1115 + op-amp.'),
    (ISL, 'Küçük titanyum / PTFE ısıtıcı', 'Asitte güvenli 50–100 W ısıtıcı.', 'AliExpress; ya da yerli rezistans üreticisinden teklif (aliserezistans, isielektrik).'),
    (CNC, 'Ön yüklü çift somun ya da C5 taşlanmış bilyalı vida', '±0,02 mm tekrarlanabilirlik.', 'Önce Mermak / Rulmansepetim\'e sor; yoksa Hiwin/TBI distribütörü.'),
]

UYARILAR = [
    'Fiyatlar 2026-10-04\'te görülen KDV dahil TL fiyatlarıdır; stok ve fiyat her gün değişir. USD bazlı satıcılar (Rulmansepetim, Meon, cnc-marketi) kurla oynar.',
    'Doğrulandı = "hayır" satırları (Hepsiburada, Trendyol, n11, Amazon.com.tr, Akakçe, Koçtaş vb.) yalnız arama özetinden; siparişten önce sayfada kontrol et.',
    'Fiyatı boş satırlar toplamda yok sayılır (Özet sayfasında sayısı yazar): ısıtıcı plaka, konum pimleri, turuncu akrilik, atık bidonu.',
    'Kargo, gümrük, cıvata/somun, alüminyum plakalar, 3D baskı parçalar ve PCB\'ler dahil değil.',
    'Kimyasallar: bazı satıcılar şahıslara satmıyor (Rokim/laboratuvar siteleri, Kimyacınız aseton), Kimyacınız HCl\'yi kargolamıyor. H2O2 ADR 5.1, HCl ADR 8: kargo reddedebilir. Ticari aşamada şirket faturasıyla al.',
    'Mono LCD panellerin init dizisi/arayüzü yayımlanmamış: ESP32-P4 ile sürmek deneysel. Satılan ürün için panel üreticisinden datasheet (NDA/MOQ) gerekir; yedek parçayı tersine mühendislik yalnız prototip içindir.',
    'CNC: hazır 3018 kitleri ±0,1 mm; hedef ±0,02 mm için ray + bilyalı vida (bu BOM) ve düşük salgılı mil gerekir. Kitlerle gelen GRBL GPL lisanslı: ürüne girmez.',
    'Akvaryum ısıtıcı termostatı ~32–34 °C\'de keser; 40–50 °C aşındırma için ESP32 + SSR ile sür ve aşırı sıcaklık kesicisi ekle.',
    'Bakırlı atık kanalizasyona dökülmez. HCl buharı metalleri paslandırır: ıslak istasyon CNC ve elektronikten ayrı, havalandırmalı kabinde.',
    'Robotistan sayfalarında gizli "Tükendi" etiketi var; stok, sayfanın ürün verisinden (stok adedi) okundu.',
]

HDR = ['No', 'Modül', 'Kalem', 'Özellik', 'Adet', 'Birim', 'Birim fiyat (TL, KDV dahil)', 'Toplam (TL)',
       'Satıcı', 'Link', 'Stok (%s)' % DATE, 'Doğrulandı', 'Not / alternatif']

HFONT = Font(bold=True, color='FFFFFF')
HFILL = PatternFill('solid', fgColor='1F3864')
WARN = PatternFill('solid', fgColor='FFF2CC')
THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TL = '#,##0.00 "TL"'


def header(ws, cols, widths):
    ws.append(cols)
    for i, w in enumerate(widths, 1):
        c = ws.cell(row=1, column=i)
        c.font, c.fill, c.border = HFONT, HFILL, BORDER
        c.alignment = Alignment(wrap_text=True, vertical='center')
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = 'A2'


def body_style(ws, wrap_cols):
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.border = BORDER
            c.alignment = Alignment(wrap_text=c.column in wrap_cols, vertical='top')


def link(c):
    if isinstance(c.value, str) and c.value.startswith('http'):
        c.hyperlink = c.value
        c.font = Font(color='0563C1', underline='single')


def main():
    wb = Workbook()

    # --- Ozet ---
    oz = wb.active
    oz.title = 'Özet'
    header(oz, ['Modül', 'Kalem sayısı', 'Fiyatı bilinmeyen', 'Toplam (TL, KDV dahil)'], [26, 14, 18, 24])
    mods = [POZ, CNC, ISL, HOT, ELK]
    n = len(BOM) + 1
    for i, m in enumerate(mods, 2):
        oz.append([m, '=COUNTIF(BOM!B2:B%d,A%d)' % (n, i),
                   '=COUNTIFS(BOM!B2:B%d,A%d,BOM!G2:G%d,"")' % (n, i, n),
                   '=SUMIF(BOM!B2:B%d,A%d,BOM!H2:H%d)' % (n, i, n)])
    last = len(mods) + 1
    oz.append(['TOPLAM', '=SUM(B2:B%d)' % last, '=SUM(C2:C%d)' % last, '=SUM(D2:D%d)' % last])
    body_style(oz, set())
    for r in range(2, last + 2):
        oz.cell(row=r, column=4).number_format = TL
    for c in oz[last + 1]:
        c.font = Font(bold=True)
    oz.append([])
    oz.append(['PCB üretim makinesi, prototip v0. Fiyatlar %s, Türkiye satıcıları. Kaynak: bom_v0.py' % DATE])
    oz.append(['CNC: ana liste özel yapım (ray + bilyalı vida). Hazır kitler "Alternatifler" sayfasında.'])

    # --- BOM ---
    ws = wb.create_sheet('BOM')
    header(ws, HDR, [5, 16, 26, 44, 7, 7, 14, 14, 18, 40, 18, 11, 60])
    for i, (mod, kalem, oz_, adet, birim, fiyat, sat, url, stok, dog, nt) in enumerate(BOM, 1):
        r = i + 1
        ws.append([i, mod, kalem, oz_, adet, birim, fiyat, '=IF(G%d="","",E%d*G%d)' % (r, r, r),
                   sat, url, stok, dog, nt])
    body_style(ws, {3, 4, 9, 11, 13})
    for row in ws.iter_rows(min_row=2):
        row[6].number_format = row[7].number_format = TL
        link(row[9])
        if row[11].value != 'evet' or row[6].value is None:
            row[11].fill = WARN
    ws.auto_filter.ref = 'A1:M%d' % n

    # --- Alternatifler ---
    al = wb.create_sheet('Alternatifler')
    header(al, ['Modül', 'Kalem', 'Özellik', 'Fiyat (TL, KDV dahil)', 'Satıcı', 'Link', 'Stok', 'Not'],
           [16, 28, 40, 16, 16, 40, 16, 50])
    for row in ALT:
        al.append(list(row))
    body_style(al, {2, 3, 8})
    for row in al.iter_rows(min_row=2):
        row[3].number_format = TL
        link(row[5])

    # --- Yurt disi ---
    yd = wb.create_sheet('Yurt dışı')
    header(yd, ['Modül', 'Kalem', 'Ne için', 'Nereden'], [16, 40, 46, 60])
    for row in YURTDISI:
        yd.append(list(row))
    body_style(yd, {2, 3, 4})

    # --- Uyarilar ---
    uy = wb.create_sheet('Uyarılar')
    header(uy, ['#', 'Uyarı'], [5, 120])
    for i, t in enumerate(UYARILAR, 1):
        uy.append([i, t])
    body_style(uy, {2})

    wb.save(OUT)

    tot = {}
    for r in BOM:
        if r[5] is not None:
            tot[r[0]] = tot.get(r[0], 0) + r[3] * r[5]
    for m in mods:
        print('%-18s %12.2f TL' % (m, tot.get(m, 0)))
    print('%-18s %12.2f TL  (%d kalem, %d fiyatsiz)' % ('TOPLAM', sum(tot.values()), len(BOM),
                                                        sum(1 for r in BOM if r[5] is None)))
    print('->', OUT)


if __name__ == '__main__':
    main()
